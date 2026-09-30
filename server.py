from pathlib import Path
from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
from urllib.parse import urlparse, parse_qs, urlencode
from urllib.request import urlopen, Request
from datetime import datetime, timezone, timedelta
from email.message import EmailMessage
import base64, hashlib, hmac, html, json, os, re, secrets, smtplib, sqlite3, io
from norna_extras import analyze_resume, search_jobs, search_learning

BASE = Path(__file__).resolve().parent
DB = Path(os.getenv('NORNA_DB_PATH', str(BASE / 'norna.db')))
DATABASE_URL = os.getenv('DATABASE_URL', '').strip()
try:
    import psycopg
    from psycopg.rows import dict_row
except Exception:
    psycopg = None
    dict_row = None
DEV_MODE = os.getenv('DEV_MODE', '1') == '1'
APP_BASE_URL = os.getenv('APP_BASE_URL', '').strip().rstrip('/')
COOKIE_SECURE = os.getenv('COOKIE_SECURE', '0') == '1'
OPENAI_API_KEY = os.getenv('OPENAI_API_KEY', '').strip()
OPENAI_MODEL = os.getenv('OPENAI_MODEL', 'gpt-5.6-luna').strip()



def request_base_url(handler):
    if APP_BASE_URL:
        return APP_BASE_URL
    host = (handler.headers.get('X-Forwarded-Host') or handler.headers.get('Host') or '127.0.0.1:8000').strip()
    proto = (handler.headers.get('X-Forwarded-Proto') or '').strip()
    if not proto:
        proto = 'http' if host.startswith(('127.0.0.1', 'localhost')) else 'https'
    return f'{proto}://{host}'

EMPTY_DATA = {
    'profile': {'photo': '', 'name': '', 'area': '', 'role': '', 'level': '', 'mode': '', 'location': '', 'study': '', 'opportunity': '', 'about': ''},
    'education': [], 'experience': [], 'journey': [], 'courses': [], 'skills': [], 'savedJobs': [], 'savedLearning': []
}

def db():
    if DATABASE_URL:
        if psycopg is None:
            raise RuntimeError('DATABASE_URL foi configurado, mas psycopg não está instalado.')
        return psycopg.connect(DATABASE_URL, row_factory=dict_row)
    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    return con

def sql(text):
    return text.replace('?', '%s') if DATABASE_URL else text

def execute(con, statement, params=()):
    return execute(con, sql(statement), params)

def init_db():
    con = db()
    cur = con.cursor()
    if DATABASE_URL:
        cur.execute('CREATE TABLE IF NOT EXISTS users(id BIGSERIAL PRIMARY KEY,email TEXT UNIQUE NOT NULL,password_hash TEXT NOT NULL,salt TEXT NOT NULL,verified INTEGER NOT NULL DEFAULT 0,verify_token TEXT,reset_token TEXT,created_at TEXT NOT NULL)')
        cur.execute('CREATE TABLE IF NOT EXISTS sessions(token TEXT PRIMARY KEY,user_id BIGINT NOT NULL REFERENCES users(id) ON DELETE CASCADE,expires_at TEXT NOT NULL)')
        cur.execute('CREATE TABLE IF NOT EXISTS user_data(user_id BIGINT PRIMARY KEY REFERENCES users(id) ON DELETE CASCADE,data_json TEXT NOT NULL,updated_at TEXT NOT NULL)')
    else:
        cur.execute('CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY AUTOINCREMENT,email TEXT UNIQUE NOT NULL,password_hash TEXT NOT NULL,salt TEXT NOT NULL,verified INTEGER NOT NULL DEFAULT 0,verify_token TEXT,reset_token TEXT,created_at TEXT NOT NULL)')
        cur.execute('CREATE TABLE IF NOT EXISTS sessions(token TEXT PRIMARY KEY,user_id INTEGER NOT NULL,expires_at TEXT NOT NULL,FOREIGN KEY(user_id) REFERENCES users(id))')
        cur.execute('CREATE TABLE IF NOT EXISTS user_data(user_id INTEGER PRIMARY KEY,data_json TEXT NOT NULL,updated_at TEXT NOT NULL,FOREIGN KEY(user_id) REFERENCES users(id))')
    con.commit()
    con.close()

def hash_password(password, salt=None):
    salt_bytes = base64.b64decode(salt) if salt else secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac('sha256', password.encode(), salt_bytes, 240000)
    return base64.b64encode(digest).decode(), base64.b64encode(salt_bytes).decode()

def verify_password(password, stored, salt):
    got, _ = hash_password(password, salt)
    return hmac.compare_digest(got, stored)

def valid_email(email):
    return re.match(r'^[^@\s]+@[^@\s]+\.[^@\s]+$', email or '') is not None

def cookie_value(header, name):
    if not header:
        return None
    for part in header.split(';'):
        if '=' in part:
            k, v = part.strip().split('=', 1)
            if k == name:
                return v
    return None

def user_from_handler(handler):
    token = cookie_value(handler.headers.get('Cookie'), 'norna_session')
    if not token:
        return None, None
    con = db()
    row = execute(con, 'SELECT u.*,s.expires_at FROM sessions s JOIN users u ON u.id=s.user_id WHERE s.token=?', (token,)).fetchone()
    if not row:
        con.close()
        return None, token
    expires_value = row['expires_at']
    expires_dt = expires_value if isinstance(expires_value, datetime) else datetime.fromisoformat(str(expires_value))
    if expires_dt.tzinfo is None:
        expires_dt = expires_dt.replace(tzinfo=timezone.utc)
    if expires_dt < datetime.now(timezone.utc):
        execute(con, 'DELETE FROM sessions WHERE token=?', (token,))
        con.commit()
        con.close()
        return None, token
    con.close()
    return dict(row), token

def send_email(to, subject, body):
    host = os.getenv('SMTP_HOST', '').strip()
    user = os.getenv('SMTP_USER', '').strip()
    password = os.getenv('SMTP_PASSWORD', '').strip()
    sender = os.getenv('SMTP_FROM', user).strip()
    if not host or not sender:
        return False
    port = int(os.getenv('SMTP_PORT', '587'))
    use_tls = os.getenv('SMTP_TLS', '1') == '1'
    msg = EmailMessage()
    msg['From'] = sender
    msg['To'] = to
    msg['Subject'] = subject
    msg.set_content(body)
    try:
        with smtplib.SMTP(host, port, timeout=15) as smtp:
            if use_tls:
                smtp.starttls()
            if user:
                smtp.login(user, password)
            smtp.send_message(msg)
        return True
    except Exception as exc:
        print('Falha SMTP:', exc)
        return False

def strip_html(text):
    text = re.sub(r'<[^>]+>', ' ', text or '')
    return re.sub(r'\s+', ' ', html.unescape(text)).strip()

def parse_date(value):
    if not value:
        return None
    try:
        dt = datetime.fromisoformat(value.replace('Z', '+00:00'))
        return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)
    except Exception:
        return None

def read_json(handler):
    try:
        n = int(handler.headers.get('Content-Length', '0'))
        raw = handler.rfile.read(n)
        return json.loads(raw.decode('utf-8') or '{}')
    except Exception:
        return {}

def openai_text(prompt):
    if not OPENAI_API_KEY:
        return None
    payload = json.dumps({'model': OPENAI_MODEL, 'input': prompt}).encode('utf-8')
    req = Request(
        'https://api.openai.com/v1/responses',
        data=payload,
        headers={'Authorization': f'Bearer {OPENAI_API_KEY}', 'Content-Type': 'application/json'},
        method='POST'
    )
    try:
        with urlopen(req, timeout=45) as response:
            obj = json.loads(response.read().decode('utf-8'))
        pieces = []
        for item in obj.get('output', []):
            for content in item.get('content', []):
                if content.get('type') in ('output_text', 'text') and content.get('text'):
                    pieces.append(content['text'])
        return '\n'.join(pieces).strip() or None
    except Exception as exc:
        print('Falha IA:', exc)
        return None

def extract_uploaded_resume_text(filename, content_b64):
    raw = base64.b64decode(content_b64 or '')
    lower = (filename or '').lower()
    if lower.endswith('.pdf'):
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(raw))
            return '\n'.join((page.extract_text() or '') for page in reader.pages)
        except Exception as exc:
            raise ValueError(f'Não consegui ler este PDF: {exc}')
    if lower.endswith('.docx'):
        try:
            from docx import Document
            doc = Document(io.BytesIO(raw))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(' | '.join(cell.text.strip() for cell in row.cells if cell.text.strip()))
            return '\n'.join(parts)
        except Exception as exc:
            raise ValueError(f'Não consegui ler este Word: {exc}')
    if lower.endswith('.txt') or lower.endswith('.md'):
        return raw.decode('utf-8', errors='ignore')
    raise ValueError('Formato não suportado. Use PDF, DOCX, TXT ou MD.')

def local_coach(data, question):
    profile = data.get('profile', {})
    skills = data.get('skills', [])
    courses = data.get('courses', [])
    jobs = data.get('savedJobs', [])
    strong = [s.get('name') for s in skills if int(s.get('level', 0) or 0) >= 70]
    developing = [s.get('name') for s in skills if 0 < int(s.get('level', 0) or 0) < 70]
    lines = []
    if profile.get('role') or profile.get('area'):
        lines.append(f"Seu foco atual é {profile.get('role') or profile.get('area')}.")
    if strong:
        lines.append('Pontos fortes registrados: ' + ', '.join(strong[:5]) + '.')
    if developing:
        lines.append('Competências em desenvolvimento: ' + ', '.join(developing[:5]) + '.')
    if not skills:
        lines.append('O melhor próximo passo é cadastrar suas competências para a NORNA comparar oportunidades com mais precisão.')
    elif not courses:
        lines.append('Você já tem competências registradas; agora vale ligar um curso ou etapa da Jornada ao que deseja desenvolver.')
    elif not jobs:
        lines.append('Seu perfil já tem base suficiente para testar a aba Vagas e comparar requisitos reais.')
    else:
        lines.append('Use as lacunas que mais se repetem nas vagas salvas como prioridade de estudo.')
    return '\n'.join(lines)

class Handler(SimpleHTTPRequestHandler):
    server_version = 'NORNA/1.0'

    def log_message(self, fmt, *args):
        print('[NORNA]', fmt % args)

    def end_headers(self):
        if DEV_MODE:
            self.send_header('Cache-Control', 'no-store, no-cache, must-revalidate, max-age=0')
            self.send_header('Pragma', 'no-cache')
            self.send_header('Expires', '0')
        super().end_headers()

    def translate_path(self, path):
        clean = urlparse(path).path.lstrip('/') or 'index.html'
        target = (BASE / clean).resolve()
        if BASE.resolve() not in target.parents and target != BASE.resolve():
            return str(BASE / 'index.html')
        return str(target)

    def send_json(self, obj, status=200, headers=None):
        body = json.dumps(obj, ensure_ascii=False).encode('utf-8')
        self.send_response(status)
        self.send_header('Content-Type', 'application/json; charset=utf-8')
        self.send_header('Cache-Control', 'no-store')
        for k, v in (headers or {}).items():
            self.send_header(k, v)
        self.send_header('Content-Length', str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def redirect(self, url, headers=None):
        self.send_response(302)
        self.send_header('Location', url)
        for k, v in (headers or {}).items():
            self.send_header(k, v)
        self.end_headers()

    def require_user(self):
        user, _ = user_from_handler(self)
        if not user:
            self.send_json({'error': 'Sessão expirada. Entre novamente.'}, 401)
            return None
        return user

    def do_GET(self):
        parsed = urlparse(self.path)
        path = parsed.path
        qs = parse_qs(parsed.query)

        if path == '/api/health':
            return self.send_json({'ok': True})

        if path == '/api/me':
            user, _ = user_from_handler(self)
            if not user:
                return self.send_json({'error': 'Não autenticado'}, 401)
            return self.send_json({'id': user['id'], 'email': user['email']})

        if path == '/api/data':
            user = self.require_user()
            if not user:
                return
            con = db()
            row = execute(con, 'SELECT data_json FROM user_data WHERE user_id=?', (user['id'],)).fetchone()
            con.close()
            return self.send_json({'data': json.loads(row['data_json']) if row else EMPTY_DATA})

        if path == '/api/verify':
            token = (qs.get('token', [''])[0] or '').strip()
            con = db()
            row = execute(con, 'SELECT id FROM users WHERE verify_token=?', (token,)).fetchone()
            if row:
                execute(con, 'UPDATE users SET verified=1,verify_token=NULL WHERE id=?', (row['id'],))
                con.commit()
                con.close()
                return self.redirect('/acesso.html?verified=1')
            con.close()
            return self.redirect('/acesso.html?verified=0')

        if path == '/api/jobs':
            user = self.require_user()
            if not user:
                return
            query = (qs.get('query', [''])[0] or '').strip()
            keyword = (qs.get('keyword', [''])[0] or '').strip()
            role = (qs.get('role', [''])[0] or '').strip()
            area = (qs.get('area', [''])[0] or '').strip()
            mode = (qs.get('mode', [''])[0] or '').strip()
            days = max(1, min(60, int(qs.get('days', ['7'])[0] or 7)))
            result = search_jobs(query=query, keyword=keyword, role=role, area=area, mode=mode, days=days)
            return self.send_json(result)

        if path == '/api/learning':
            user = self.require_user()
            if not user:
                return
            kind = (qs.get('type', [''])[0] or '').strip()
            mode = (qs.get('mode', [''])[0] or '').strip()
            keyword = (qs.get('keyword', [''])[0] or '').strip()
            price = (qs.get('price', [''])[0] or '').strip()
            return self.send_json(search_learning(kind=kind, mode=mode, keyword=keyword, price=price))

        return super().do_GET()

    def do_PUT(self):
        if self.path != '/api/data':
            return self.send_json({'error': 'Rota não encontrada'}, 404)
        user = self.require_user()
        if not user:
            return
        payload = read_json(self)
        con = db()
        execute(con, 
            'INSERT INTO user_data(user_id,data_json,updated_at) VALUES(?,?,?) ON CONFLICT(user_id) DO UPDATE SET data_json=excluded.data_json,updated_at=excluded.updated_at',
            (user['id'], json.dumps(payload, ensure_ascii=False), datetime.now(timezone.utc).isoformat())
        )
        con.commit()
        con.close()
        return self.send_json({'ok': True})

    def do_POST(self):
        path = urlparse(self.path).path
        body = read_json(self)

        if path == '/api/register':
            email = (body.get('email') or '').strip().lower()
            password = body.get('password') or ''
            if not valid_email(email):
                return self.send_json({'error': 'E-mail inválido.'}, 400)
            if len(password) < 8:
                return self.send_json({'error': 'A senha precisa ter pelo menos 8 caracteres.'}, 400)
            password_hash, salt = hash_password(password)
            verify_token = secrets.token_urlsafe(32)
            con = db()
            try:
                if DATABASE_URL:
                    cur = execute(con,
                        'INSERT INTO users(email,password_hash,salt,verified,verify_token,created_at) VALUES(?,?,?,?,?,?) RETURNING id',
                        (email, password_hash, salt, 0, verify_token, datetime.now(timezone.utc).isoformat())
                    )
                    uid = cur.fetchone()['id']
                else:
                    cur = execute(con,
                        'INSERT INTO users(email,password_hash,salt,verified,verify_token,created_at) VALUES(?,?,?,?,?,?)',
                        (email, password_hash, salt, 0, verify_token, datetime.now(timezone.utc).isoformat())
                    )
                    uid = cur.lastrowid
                execute(con, 'INSERT INTO user_data(user_id,data_json,updated_at) VALUES(?,?,?)', (uid, json.dumps(EMPTY_DATA), datetime.now(timezone.utc).isoformat()))
                con.commit()
            except Exception as exc:
                con.rollback()
                con.close()
                if isinstance(exc, sqlite3.IntegrityError) or (psycopg is not None and isinstance(exc, psycopg.errors.UniqueViolation)):
                    return self.send_json({'error': 'Já existe uma conta com este e-mail.'}, 409)
                print('Falha ao criar conta:', exc)
                return self.send_json({'error': 'Não foi possível criar a conta agora.'}, 500)
            con.close()
            verify_url = f'{request_base_url(self)}/api/verify?token={verify_token}'
            if DEV_MODE:
                # Ambiente local de teste: nao depende de e-mail/SMTP.
                con = db()
                execute(con, 'UPDATE users SET verified=1,verify_token=NULL WHERE id=?', (uid,))
                con.commit()
                con.close()
                return self.send_json({'ok': True, 'dev_auto_verified': True}, 201)

            sent = send_email(email, 'Confirme sua conta NORNA', f'Bem-vindo(a) à NORNA.\n\nConfirme sua conta:\n{verify_url}\n')
            return self.send_json({'ok': True, 'email_sent': sent}, 201)

        if path == '/api/login':
            email = (body.get('email') or '').strip().lower()
            password = body.get('password') or ''
            con = db()
            user = execute(con, 'SELECT * FROM users WHERE email=?', (email,)).fetchone()
            if not user or not verify_password(password, user['password_hash'], user['salt']):
                con.close()
                return self.send_json({'error': 'E-mail ou senha incorretos.'}, 401)
            if not user['verified']:
                con.close()
                return self.send_json({'error': 'Confirme seu e-mail antes de entrar.', 'needs_verification': True}, 403)
            token = secrets.token_urlsafe(40)
            expires = datetime.now(timezone.utc) + timedelta(days=30)
            execute(con, 'DELETE FROM sessions WHERE user_id=?', (user['id'],))
            execute(con, 'INSERT INTO sessions(token,user_id,expires_at) VALUES(?,?,?)', (token, user['id'], expires.isoformat()))
            con.commit()
            con.close()
            cookie = f'norna_session={token}; Path=/; HttpOnly; SameSite=Lax; Max-Age={30*24*3600}' + ('; Secure' if COOKIE_SECURE else '')
            return self.send_json({'ok': True}, headers={'Set-Cookie': cookie})

        if path == '/api/logout':
            _, token = user_from_handler(self)
            con = db()
            if token:
                execute(con, 'DELETE FROM sessions WHERE token=?', (token,))
                con.commit()
            con.close()
            return self.send_json({'ok': True}, headers={'Set-Cookie': 'norna_session=; Path=/; HttpOnly; SameSite=Lax; Max-Age=0'})

        if path == '/api/forgot':
            email = (body.get('email') or '').strip().lower()
            con = db()
            user = execute(con, 'SELECT id FROM users WHERE email=?', (email,)).fetchone()
            result = {'message': 'Se essa conta existir, enviaremos um link de recuperação.'}
            if user:
                token = secrets.token_urlsafe(32)
                execute(con, 'UPDATE users SET reset_token=? WHERE id=?', (token, user['id']))
                con.commit()
                reset_url = f'{request_base_url(self)}/acesso.html?reset={token}'
                sent = send_email(email, 'Recupere sua senha NORNA', f'Use este link para criar uma nova senha:\n{reset_url}\n')
                if DEV_MODE and not sent:
                    result['dev_reset_url'] = reset_url
            con.close()
            return self.send_json(result)

        if path == '/api/reset':
            token = (body.get('token') or '').strip()
            password = body.get('password') or ''
            if len(password) < 8:
                return self.send_json({'error': 'A senha precisa ter pelo menos 8 caracteres.'}, 400)
            con = db()
            user = execute(con, 'SELECT id FROM users WHERE reset_token=?', (token,)).fetchone()
            if not user:
                con.close()
                return self.send_json({'error': 'Link inválido ou já utilizado.'}, 400)
            password_hash, salt = hash_password(password)
            execute(con, 'UPDATE users SET password_hash=?,salt=?,reset_token=NULL WHERE id=?', (password_hash, salt, user['id']))
            execute(con, 'DELETE FROM sessions WHERE user_id=?', (user['id'],))
            con.commit()
            con.close()
            return self.send_json({'ok': True})

        if path == '/api/resume/upload':
            user = self.require_user()
            if not user:
                return
            filename = (body.get('filename') or '').strip()
            content_b64 = body.get('content_base64') or ''
            if not filename or not content_b64:
                return self.send_json({'error': 'Envie um arquivo de currículo.'}, 400)
            try:
                text = extract_uploaded_resume_text(filename, content_b64)
            except ValueError as exc:
                return self.send_json({'error': str(exc)}, 400)
            if len(text.strip()) < 20:
                return self.send_json({'error': 'Não consegui extrair texto suficiente desse arquivo.'}, 400)
            return self.send_json({'ok': True, 'text': text[:120000], 'extracted': analyze_resume(text[:120000])})

        if path == '/api/resume/analyze':
            user = self.require_user()
            if not user:
                return
            text = (body.get('text') or '').strip()
            if len(text) < 20:
                return self.send_json({'error': 'Não consegui encontrar texto suficiente nesse currículo.'}, 400)
            if len(text) > 120000:
                text = text[:120000]
            return self.send_json({'ok': True, 'extracted': analyze_resume(text)})

        if path == '/api/ai/coach':
            user = self.require_user()
            if not user:
                return
            question = (body.get('question') or '').strip()
            if not question:
                return self.send_json({'error': 'Escreva uma pergunta.'}, 400)
            con = db()
            row = execute(con, 'SELECT data_json FROM user_data WHERE user_id=?', (user['id'],)).fetchone()
            con.close()
            data = json.loads(row['data_json']) if row else EMPTY_DATA
            prompt = (
                'Você é a NORNA, uma assistente de carreira e aprendizado. Responda em português do Brasil, de forma curta, concreta e sem prometer contratação. '
                'Use apenas os dados fornecidos. Diferencie compatibilidade de vaga de probabilidade de contratação.\n\nDADOS DO USUÁRIO:\n'
                + json.dumps(data, ensure_ascii=False)[:22000]
                + '\n\nPERGUNTA:\n' + question
            )
            answer = openai_text(prompt) or local_coach(data, question)
            return self.send_json({'answer': answer, 'ai_enabled': bool(OPENAI_API_KEY)})

        return self.send_json({'error': 'Rota não encontrada'}, 404)

if __name__ == '__main__':
    init_db()
    os.chdir(BASE)
    port = int(os.getenv('PORT', '8000'))
    host = '0.0.0.0'
    print(f'NORNA ouvindo em {host}:{port}')
    ThreadingHTTPServer((host, port), Handler).serve_forever()
